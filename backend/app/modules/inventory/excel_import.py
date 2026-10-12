"""Excel and CSV Bulk Import Engine for Hardware Catalogues.

Per Module 2 of Development Plan:
1. Preview and error report before committing.
2. Creates Categories, Brands, Products, Packaging Units, and Price Lists.
3. Initializes opening stock in base units with opening Weighted Average Cost.
"""

import csv
from datetime import datetime, timezone
import io
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, Iterable, List, Optional, Tuple
import uuid
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.catalogue import Brand, Category, PriceList, Product, ProductPrice, ProductUnit
from app.models.stock import Godown, StockBalance, StockMovement
from app.modules.inventory.unit_converter import SmartUnitConverter


class ImportErrorItem(BaseModel):
    row_number: int
    field: str
    message: str


class ParsedCatalogueRow(BaseModel):
    row_number: int
    code: str
    name: str
    tamil_name: Optional[str] = None
    category_name: str
    brand_name: Optional[str] = None
    hsn_code: str
    base_unit: str
    is_variable_weight: bool = False
    purchase_price: Decimal
    price_retail: Decimal
    price_contractor: Optional[Decimal] = None
    price_wholesale: Optional[Decimal] = None
    packing_unit: Optional[str] = None
    packing_factor: Optional[Decimal] = None
    opening_stock_base: Decimal = Decimal("0.0000")
    gst_rate_percent: Decimal = Decimal("18.00")


class ImportRowResult(BaseModel):
    row_number: int
    code: Optional[str] = None
    status: str  # 'CREATED', 'UPDATED', 'VALID' (dry run), 'ERROR'
    errors: List[ImportErrorItem] = []


class ImportResultSchema(BaseModel):
    file_name: str
    dry_run: bool
    committed: bool
    total_rows: int
    valid_rows_count: int
    invalid_rows_count: int
    products_created: int = 0
    products_updated: int = 0
    units_created: int = 0
    prices_created: int = 0
    rows: List[ImportRowResult]


class ImportPreviewResult(BaseModel):
    total_rows: int
    valid_rows_count: int
    invalid_rows_count: int
    errors: List[ImportErrorItem]
    sample_valid_rows: List[ParsedCatalogueRow]


def _cell_text(value: Any) -> str:
    """Normalise a CSV / Excel cell to text (Excel gives 7318.0 for a numeric HSN cell)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Y" if value else "N"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value)


class CatalogueImportEngine:
    VALID_BASE_UNITS = {"piece", "kg", "metre", "sqft", "litre"}
    VALID_GST_RATES = {Decimal("0"), Decimal("5"), Decimal("12"), Decimal("18"), Decimal("28")}
    MAX_ROWS = 10000

    @classmethod
    def parse_csv_stream(cls, text_content: str) -> Tuple[List[ParsedCatalogueRow], List[ImportErrorItem]]:
        """Parse raw CSV string into validated rows and error list."""
        reader = csv.DictReader(io.StringIO(text_content))
        return cls.parse_records(enumerate(reader, start=2))  # row 1 is header

    @classmethod
    def parse_csv_bytes(cls, content: bytes) -> Tuple[List[ParsedCatalogueRow], List[ImportErrorItem]]:
        try:
            text_content = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("CSV must be UTF-8 encoded (save as 'CSV UTF-8' in Excel)")
        return cls.parse_csv_stream(text_content)

    @classmethod
    def parse_xlsx_bytes(cls, content: bytes) -> Tuple[List[ParsedCatalogueRow], List[ImportErrorItem]]:
        """Parse the first worksheet of an .xlsx file; row 1 holds the column headers."""
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise ValueError("Could not read the Excel file. Upload a valid .xlsx workbook.")
        try:
            ws = wb.worksheets[0]
            rows_iter = ws.iter_rows(values_only=True)
            header = next(rows_iter, None)
            if not header:
                return [], []
            keys = [_cell_text(h).strip().lower() for h in header]
            records = []
            for row_idx, values in enumerate(rows_iter, start=2):
                if values is None or all(v is None or _cell_text(v).strip() == "" for v in values):
                    continue  # skip blank lines
                records.append((row_idx, {k: _cell_text(v) for k, v in zip(keys, values) if k}))
        finally:
            wb.close()
        return cls.parse_records(records)

    @classmethod
    def parse_records(
        cls, records: Iterable[Tuple[int, Dict[str, Any]]]
    ) -> Tuple[List[ParsedCatalogueRow], List[ImportErrorItem]]:
        """Validate (row_number, {column: text}) records into catalogue rows and per-row errors."""
        valid_rows: List[ParsedCatalogueRow] = []
        errors: List[ImportErrorItem] = []
        seen_codes: Dict[str, int] = {}

        for count, (row_idx, raw) in enumerate(records, start=1):
            if count > cls.MAX_ROWS:
                raise ValueError(f"Too many rows (max {cls.MAX_ROWS} per file)")
            raw = {str(k).strip().lower(): ("" if v is None else str(v)) for k, v in raw.items() if k is not None}
            code = raw.get("code", "").strip()
            name = raw.get("name", "").strip()
            tamil_name = raw.get("tamil_name", "").strip() or None
            category = raw.get("category", "").strip() or "General"
            brand = raw.get("brand", "").strip() or None
            hsn = raw.get("hsn_code", "").strip()
            base_unit = raw.get("base_unit", "").strip().lower()
            var_weight_str = raw.get("is_variable_weight", "N").strip().upper()
            is_var_weight = var_weight_str in ("Y", "YES", "TRUE", "1")

            # Validations
            if not code:
                errors.append(ImportErrorItem(row_number=row_idx, field="code", message="Product code is required"))
                continue
            if not name:
                errors.append(ImportErrorItem(row_number=row_idx, field="name", message="Product name is required"))
                continue
            if not hsn or not (4 <= len(hsn) <= 8) or not hsn.isdigit():
                errors.append(ImportErrorItem(row_number=row_idx, field="hsn_code", message="HSN must be 4 to 8 digits"))
                continue
            if base_unit not in cls.VALID_BASE_UNITS:
                errors.append(ImportErrorItem(row_number=row_idx, field="base_unit", message=f"Base unit must be one of {cls.VALID_BASE_UNITS}"))
                continue
            if code.upper() in seen_codes:
                errors.append(ImportErrorItem(
                    row_number=row_idx, field="code",
                    message=f"Duplicate code {code} (already on row {seen_codes[code.upper()]})",
                ))
                continue

            gst_text = raw.get("gst_rate_percent", "").strip().rstrip("%")
            try:
                gst_rate = Decimal(gst_text) if gst_text else Decimal("18")
            except InvalidOperation:
                gst_rate = None
            if gst_rate is None or gst_rate not in cls.VALID_GST_RATES:
                errors.append(ImportErrorItem(row_number=row_idx, field="gst_rate_percent", message="GST rate must be 0, 5, 12, 18 or 28"))
                continue

            try:
                purchase_price = Decimal(raw.get("purchase_price", "0").strip() or "0")
                price_retail = Decimal(raw.get("selling_price_retail", "0").strip() or "0")
                price_contractor = Decimal(raw.get("selling_price_contractor", "0").strip() or "0") if raw.get("selling_price_contractor") else None
                price_wholesale = Decimal(raw.get("selling_price_wholesale", "0").strip() or "0") if raw.get("selling_price_wholesale") else None
                opening_stock = Decimal(raw.get("opening_stock_base", "0").strip() or "0")
            except InvalidOperation:
                errors.append(ImportErrorItem(row_number=row_idx, field="pricing", message="Invalid numeric format for prices or opening stock"))
                continue

            # Packaging unit (e.g. Box = 100)
            packing_unit = raw.get("packing_unit", "").strip() or None
            packing_factor = None
            if packing_unit:
                try:
                    packing_factor = Decimal(raw.get("packing_factor", "1").strip() or "1")
                except InvalidOperation:
                    errors.append(ImportErrorItem(row_number=row_idx, field="packing_factor", message="Invalid packing factor"))
                    continue

            valid_rows.append(
                ParsedCatalogueRow(
                    row_number=row_idx,
                    code=code,
                    name=name,
                    tamil_name=tamil_name,
                    category_name=category,
                    brand_name=brand,
                    hsn_code=hsn,
                    base_unit=base_unit,
                    is_variable_weight=is_var_weight,
                    purchase_price=purchase_price,
                    price_retail=price_retail,
                    price_contractor=price_contractor,
                    price_wholesale=price_wholesale,
                    packing_unit=packing_unit,
                    packing_factor=packing_factor,
                    opening_stock_base=opening_stock,
                    gst_rate_percent=gst_rate.quantize(Decimal("0.01")),
                )
            )
            seen_codes[code.upper()] = row_idx

        return valid_rows, errors

    @classmethod
    async def import_rows(
        cls,
        tenant_id: str,
        file_name: str,
        valid_rows: List[ParsedCatalogueRow],
        errors: List[ImportErrorItem],
        session: AsyncSession,
        dry_run: bool = False,
    ) -> ImportResultSchema:
        """Commit valid rows (unless dry_run) and build the per-row result summary."""
        row_results: Dict[int, ImportRowResult] = {}
        for err in errors:
            res = row_results.setdefault(err.row_number, ImportRowResult(row_number=err.row_number, status="ERROR"))
            res.errors.append(err)

        summary: Dict[str, Any] = {}
        if valid_rows and not dry_run:
            summary = await cls.commit_catalogue(tenant_id, valid_rows, session)
            created_codes = set(summary.get("created_codes", []))
            for r in valid_rows:
                row_results[r.row_number] = ImportRowResult(
                    row_number=r.row_number, code=r.code,
                    status="CREATED" if r.code.upper() in created_codes else "UPDATED",
                )
        else:
            for r in valid_rows:
                row_results[r.row_number] = ImportRowResult(row_number=r.row_number, code=r.code, status="VALID")

        return ImportResultSchema(
            file_name=file_name,
            dry_run=dry_run,
            committed=bool(valid_rows) and not dry_run,
            total_rows=len(row_results),
            valid_rows_count=len(valid_rows),
            invalid_rows_count=len({e.row_number for e in errors}),
            products_created=summary.get("products_created", 0),
            products_updated=summary.get("products_updated", 0),
            units_created=summary.get("units_created", 0),
            prices_created=summary.get("prices_created", 0),
            rows=[row_results[k] for k in sorted(row_results)],
        )

    @classmethod
    async def commit_catalogue(
        cls,
        tenant_id: str,
        rows: List[ParsedCatalogueRow],
        session: AsyncSession,
    ) -> Dict[str, Any]:
        """Atomically commit validated rows to the database."""
        # Pre-load or create categories
        categories: Dict[str, str] = {}
        cat_result = await session.execute(select(Category).where(Category.tenant_id == tenant_id))
        for c in cat_result.scalars():
            categories[c.name.lower()] = c.id

        # Pre-load or create brands
        brands: Dict[str, str] = {}
        brand_result = await session.execute(select(Brand).where(Brand.tenant_id == tenant_id))
        for b in brand_result.scalars():
            brands[b.name.lower()] = b.id

        # Pre-load price lists
        price_lists: Dict[str, str] = {}
        pl_result = await session.execute(select(PriceList).where(PriceList.tenant_id == tenant_id))
        for pl in pl_result.scalars():
            price_lists[pl.code.upper()] = pl.id

        # Get primary godown
        godown_result = await session.execute(
            select(Godown).where(Godown.tenant_id == tenant_id, Godown.is_primary.is_(True))
        )
        primary_godown = godown_result.scalar_one_or_none()
        if not primary_godown:
            primary_godown = (
                await session.execute(select(Godown).where(Godown.tenant_id == tenant_id).limit(1))
            ).scalar_one_or_none()
        if not primary_godown and any(r.opening_stock_base > Decimal("0") for r in rows):
            raise ValueError("No godown configured for tenant; cannot post opening stock")
        primary_godown_id = primary_godown.id if primary_godown else None

        # Pre-load existing products by code
        existing_products: Dict[str, Product] = {}
        prod_result = await session.execute(
            select(Product).where(Product.tenant_id == tenant_id)
        )
        for p in prod_result.scalars():
            existing_products[p.code.upper()] = p

        products_created = 0
        products_updated = 0
        units_created = 0
        prices_created = 0
        created_codes: List[str] = []

        for r in rows:
            # Category
            cat_id = categories.get(r.category_name.lower())
            if not cat_id:
                cat_id = str(uuid.uuid4())
                session.add(Category(id=cat_id, tenant_id=tenant_id, name=r.category_name))
                categories[r.category_name.lower()] = cat_id

            # Brand
            brand_id = None
            if r.brand_name:
                brand_id = brands.get(r.brand_name.lower())
                if not brand_id:
                    brand_id = str(uuid.uuid4())
                    session.add(Brand(id=brand_id, tenant_id=tenant_id, name=r.brand_name))
                    brands[r.brand_name.lower()] = brand_id

            # Check if product exists
            existing_prod = existing_products.get(r.code.upper())
            if existing_prod:
                # Update existing
                existing_prod.name = r.name
                existing_prod.tamil_name = r.tamil_name
                existing_prod.category_id = cat_id
                existing_prod.brand_id = brand_id
                existing_prod.hsn_code = r.hsn_code
                existing_prod.base_unit = r.base_unit
                existing_prod.is_variable_weight = r.is_variable_weight
                existing_prod.purchase_price_base = r.purchase_price
                existing_prod.gst_rate_percent = r.gst_rate_percent
                product_id = existing_prod.id
                products_updated += 1
            else:
                # Product insert
                product_id = str(uuid.uuid4())
                product = Product(
                    id=product_id,
                    tenant_id=tenant_id,
                    code=r.code,
                    name=r.name,
                    tamil_name=r.tamil_name,
                    category_id=cat_id,
                    brand_id=brand_id,
                    hsn_code=r.hsn_code,
                    base_unit=r.base_unit,
                    is_variable_weight=r.is_variable_weight,
                    purchase_price_base=r.purchase_price,
                    gst_rate_percent=r.gst_rate_percent,
                )
                session.add(product)
                products_created += 1
                created_codes.append(r.code.upper())

                # Base Unit
                base_unit_obj = ProductUnit(
                    id=str(uuid.uuid4()),
                    tenant_id=tenant_id,
                    product_id=product_id,
                    unit_name=r.base_unit,
                    conversion_factor=Decimal("1.0"),
                    is_default_sales_unit=(r.packing_unit is None),
                    is_default_purchase_unit=False,
                )
                session.add(base_unit_obj)
                units_created += 1

            # Packaging Unit if defined (skip if a re-imported product already has it)
            existing_unit_names = (
                {u.unit_name.lower() for u in existing_prod.units} if existing_prod else set()
            )
            if r.packing_unit and r.packing_factor and r.packing_unit.lower() not in existing_unit_names:
                pack_unit_obj = ProductUnit(
                    id=str(uuid.uuid4()),
                    tenant_id=tenant_id,
                    product_id=product_id,
                    unit_name=r.packing_unit,
                    conversion_factor=r.packing_factor,
                    is_default_sales_unit=True,
                    is_default_purchase_unit=True,
                )
                session.add(pack_unit_obj)
                units_created += 1

            # Re-import replaces current prices: archive the active rows first (history is preserved)
            if existing_prod:
                for old_price in existing_prod.prices:
                    if old_price.effective_to is None and old_price.unit_name == r.base_unit:
                        old_price.effective_to = datetime.now(timezone.utc)

            # Retail Price
            if "RETAIL" in price_lists:
                session.add(
                    ProductPrice(
                        id=str(uuid.uuid4()),
                        tenant_id=tenant_id,
                        product_id=product_id,
                        unit_name=r.base_unit,
                        price_list_id=price_lists["RETAIL"],
                        price=r.price_retail,
                    )
                )
                prices_created += 1

            # Contractor Price
            if r.price_contractor and "CONTRACTOR" in price_lists:
                session.add(
                    ProductPrice(
                        id=str(uuid.uuid4()),
                        tenant_id=tenant_id,
                        product_id=product_id,
                        unit_name=r.base_unit,
                        price_list_id=price_lists["CONTRACTOR"],
                        price=r.price_contractor,
                    )
                )
                prices_created += 1

            # Wholesale Price
            if r.price_wholesale and "WHOLESALE" in price_lists:
                session.add(
                    ProductPrice(
                        id=str(uuid.uuid4()),
                        tenant_id=tenant_id,
                        product_id=product_id,
                        unit_name=r.base_unit,
                        price_list_id=price_lists["WHOLESALE"],
                        price=r.price_wholesale,
                    )
                )
                prices_created += 1

            # Opening Stock
            if r.opening_stock_base > Decimal("0"):
                session.add(
                    StockMovement(
                        id=str(uuid.uuid4()),
                        tenant_id=tenant_id,
                        godown_id=primary_godown_id,
                        product_id=product_id,
                        movement_type="OPENING",
                        quantity_base=r.opening_stock_base,
                        unit_cost_base=r.purchase_price,
                        reference_doc_type="OPENING_STOCK",
                        reference_doc_id=product_id,
                        reference_doc_number="OPENING",
                    )
                )

                bal_res = await session.execute(
                    select(StockBalance).where(
                        StockBalance.tenant_id == tenant_id,
                        StockBalance.godown_id == primary_godown_id,
                        StockBalance.product_id == product_id,
                    )
                )
                existing_bal = bal_res.scalar_one_or_none()
                if existing_bal:
                    # Balance must equal the sum of movements: add the inward OPENING qty and re-average cost
                    existing_bal.weighted_avg_cost_base = SmartUnitConverter.calculate_weighted_average_cost(
                        current_base_qty=existing_bal.current_quantity_base,
                        current_wac_base=existing_bal.weighted_avg_cost_base,
                        inward_base_qty=r.opening_stock_base,
                        inward_landed_cost_base=r.purchase_price,
                    )
                    existing_bal.current_quantity_base += r.opening_stock_base
                else:
                    session.add(
                        StockBalance(
                            id=str(uuid.uuid4()),
                            tenant_id=tenant_id,
                            godown_id=primary_godown_id,
                            product_id=product_id,
                            current_quantity_base=r.opening_stock_base,
                            weighted_avg_cost_base=r.purchase_price,
                        )
                    )

        await session.commit()

        return {
            "status": "success",
            "products_created": products_created,
            "products_updated": products_updated,
            "units_created": units_created,
            "prices_created": prices_created,
            "created_codes": created_codes,
        }
